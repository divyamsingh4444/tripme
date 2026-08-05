"""
Shared Word-document helpers for the TripME Phase 2 deliverables.

The documents are generated from a script rather than hand-edited so that a
correction made once — a persona name, an agent count, a claim about what the
product actually does — propagates everywhere it appears. The single biggest
defect in the drafts these replace was the same fact stated three different
ways in three sections.
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PINE = RGBColor(0x14, 0x45, 0x3D)
INK = RGBColor(0x0E, 0x1F, 0x1A)
SLATE = RGBColor(0x5B, 0x6B, 0x66)
FLAG = RGBColor(0xB4, 0x36, 0x1E)


def new_doc(title: str, subtitle: str, strapline: str) -> Document:
    doc = Document()

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, colour, before in (
        ('Heading 1', 17, PINE, 20),
        ('Heading 2', 13.5, PINE, 14),
        ('Heading 3', 11.5, INK, 10),
    ):
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.color.rgb = colour
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.85)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(90)
    r = t.add_run(title)
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = PINE

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(subtitle)
    r.font.size = Pt(13)
    r.font.color.rgb = INK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(strapline)
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = SLATE
    return doc


def page_break(doc):
    doc.add_page_break()


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def h3(doc, text):
    doc.add_heading(text, level=3)


def para(doc, text, italic=False, size=10.5, space_after=6):
    """Supports **bold** inline runs, which is most of what this needs."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    for i, chunk in enumerate(text.split('**')):
        if not chunk:
            continue
        r = p.add_run(chunk)
        r.font.size = Pt(size)
        r.font.bold = i % 2 == 1
        r.font.italic = italic
    return p


def bullets(doc, items, numbered=False):
    style = 'List Number' if numbered else 'List Bullet'
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        for i, chunk in enumerate(str(item).split('**')):
            if not chunk:
                continue
            r = p.add_run(chunk)
            r.font.size = Pt(10.5)
            r.font.bold = i % 2 == 1


def _shade(cell, hex_colour):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def table(doc, headers, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, head in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(str(head))
        r.font.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(hdr[i], '14453D')

    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            for j, chunk in enumerate(str(value).split('**')):
                if not chunk:
                    continue
                r = p.add_run(chunk)
                r.font.size = Pt(font)
                r.font.bold = j % 2 == 1

    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(doc, label, text, colour=PINE):
    """A single-cell shaded box, used for decisions and honesty notes."""
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    cell = t.rows[0].cells[0]
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f'{label} ')
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = colour
    for i, chunk in enumerate(text.split('**')):
        if not chunk:
            continue
        r = p.add_run(chunk)
        r.font.size = Pt(10)
        r.font.bold = i % 2 == 1
    _shade(cell, 'F1F5F3')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def placeholder(doc, what, why=''):
    """Marks something the team must supply, so nothing ships silently blank."""
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    cell = t.rows[0].cells[0]
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run('[ TO BE INSERTED ] ')
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = FLAG
    r = p.add_run(what)
    r.font.size = Pt(10)
    if why:
        r = p.add_run(f'  —  {why}')
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = SLATE
    _shade(cell, 'FDF3F1')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def save(doc, path):
    doc.save(path)
    print('wrote', path)
